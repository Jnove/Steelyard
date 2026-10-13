"""为离线前端导出真实引擎结果；源代码目录只读。

python -B scripts/export_frontend_fixtures.py --mcdm-root D:/大学/海之子
可选 --grid-step 10 导出四维单纯形上全部 286 个 10% 步长权重点。
归一化权重精确匹配时才可使用网格快照，不插值、不吸附到邻近点。
所有路由运行于隔离的内存数据库，不接触应用数据库。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

ROOT = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True


def source_hashes(source: Path) -> dict[str, str]:
    """文本统一换行后计算指纹，让 Windows 导出与 Linux CI 的检查一致。"""
    hashes: dict[str, str] = {}
    for base in [source, ROOT]:
        paths = [
            *(base / "backend/app").rglob("*.py"),
            *(base / "rules").glob("*.yaml"),
        ]
        for path in sorted(paths):
            hashes[path.relative_to(base).as_posix()] = hashlib.sha256(
                path.read_text(encoding="utf-8").encode("utf-8")
            ).hexdigest()
    for path in [*(ROOT / "data/schemes/mock").glob("*.json"), Path(__file__).resolve()]:
        hashes[path.relative_to(ROOT).as_posix()] = hashlib.sha256(
            path.read_text(encoding="utf-8").encode("utf-8")
        ).hexdigest()
    template = ROOT / "data/schemes/方案录入模板.xlsx"
    hashes[template.relative_to(ROOT).as_posix()] = hashlib.sha256(
        template.read_bytes()
    ).hexdigest()
    return hashes


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--mcdm-root",
        type=Path,
        default=ROOT,
        help="包含 T4 backend/app/api/rank.py 的只读仓库",
    )
    parser.add_argument(
        "--check", action="store_true", help="只核对引擎、规则与样例指纹，不改文件"
    )
    parser.add_argument(
        "--grid-step",
        type=int,
        default=0,
        help="能整除 100 且不小于 5 的百分比步长；0 表示关闭网格",
    )
    args = parser.parse_args()
    if args.grid_step and (args.grid_step < 5 or 100 % args.grid_step):
        parser.error("--grid-step 必须整除 100 且不小于 5，或为 0")
    source = args.mcdm_root.resolve()
    if not (source / "backend/app/api/rank.py").is_file():
        parser.error("缺少 T4 rank.py，请通过 --mcdm-root 指定草案仓库")
    hashes = source_hashes(source)
    out = ROOT / "frontend/src/mocks/fixtures"
    if args.check:
        fixture = json.loads((out / "data.json").read_text(encoding="utf-8"))
        recorded = fixture.get("source_hashes", {})
        changed = sorted(
            key
            for key in hashes.keys() | recorded.keys()
            if hashes.get(key) != recorded.get(key)
        )
        if changed or fixture.get("grid_step") != args.grid_step:
            print("离线快照已过期，请运行 export_frontend_fixtures.py --grid-step 10：")
            print("\n".join(changed) or "网格步长不一致")
            return 1
        print("离线评估、角色、字段与规则快照指纹一致")
        return 0
    sys.path.insert(0, str(ROOT / "backend"))
    # 公共模型、评估与字段元数据使用当前工作树；仅 T4 缺失模块从只读目录加载。
    import app.api
    import app.engine

    app.api.__path__.append(str(source / "backend/app/api"))
    app.engine.__path__.append(str(source / "backend/app/engine"))

    # evaluate/rank 会落 Agent 缓存；建表前须注册该模型，避免第一次调用时才导入。
    from app.agents.storage import AgentResultRow  # noqa: F401
    from app.api.evaluate import evaluate_project
    from app.api.projects import ProjectRead
    from app.api.rank import RankRequest, list_roles, rank_project
    from app.api.rules import build_rules_response
    from app.api.schema import build_all_schemas
    from app.api.schemes import SchemeRead
    from app.excel import parse_workbook, write_export
    from app.models import Project, ProjectParams, Scheme, SchemeParams
    from openpyxl import load_workbook
    from sqlmodel import Session, SQLModel, create_engine

    out.mkdir(parents=True, exist_ok=True)
    mock_dir = ROOT / "data/schemes/mock"
    filenames = [
        "A_高品质.json",
        "B_不合规.json",
        "C_MiC高标准化.json",
        "D_均衡.json",
        "E_被支配.json",
        "C_现浇.json",
    ]

    def payload(filename: str) -> dict:
        value = json.loads((mock_dir / filename).read_text(encoding="utf-8"))
        if "虚构数据" not in value.get("_note", ""):
            raise ValueError(f"缺少虚构数据声明：{filename}")
        return value

    fixture_id = hashlib.sha256(json.dumps(hashes, sort_keys=True).encode()).hexdigest()
    stamp = datetime(2026, 1, 1, tzinfo=UTC)
    engine = create_engine("sqlite://")
    SQLModel.metadata.create_all(engine)
    with Session(engine) as session:
        raw = payload("project.json")
        project = Project(
            id=1,
            name=raw["name"],
            params=ProjectParams.model_validate(raw["params"]).model_dump(),
            created_at=stamp,
            updated_at=stamp,
        )
        session.add(project)
        rows = []
        for i, filename in enumerate(filenames, 1):
            raw = payload(filename)
            row = Scheme(
                id=i,
                project_id=1,
                name=raw["name"],
                params=SchemeParams.model_validate(raw["params"]).model_dump(),
                is_whatif=i == 6,
                parent_id=3 if i == 6 else None,
                created_at=stamp,
                updated_at=stamp,
            )
            rows.append(row)
            session.add(row)
        session.commit()
        evaluation = evaluate_project(1, session, include_whatif=False).model_dump(
            mode="json"
        )
        roles = list_roles(1, session)

        def snapshot(request: RankRequest) -> dict:
            value = rank_project(1, request, session).model_dump(mode="json")
            value["elapsed_ms"] = (
                0  # 离线快照，不把导出耗时冒充当前浏览器或服务器耗时。
            )
            return value

        ranks = {
            r["role"]: snapshot(RankRequest(role=r["role"])) for r in roles["roles"]
        }
        grid = []
        if args.grid_step:
            units = 100 // args.grid_step
            for a in range(units + 1):
                for b in range(units - a + 1):
                    for c in range(units - a - b + 1):
                        weights = [x / units for x in (a, b, c, units - a - b - c)]
                        grid.append(snapshot(RankRequest(weights=weights)))

        project_read = ProjectRead.model_validate(project).model_dump(mode="json")
        scheme_reads = [
            SchemeRead.model_validate(row).model_dump(mode="json") for row in rows
        ]
        # 导出时调用真实 Excel 解析器；浏览器只识别内容完全相同的文件。
        workbook = out / "demo-import.xlsx"
        write_export(
            workbook,
            project.params,
            [{"name": r.name, "params": r.params} for r in rows if not r.is_whatif],
        )
        invalid_workbook = out / "demo-invalid.xlsx"
        wb = load_workbook(workbook)
        ws = wb["方案参数"]
        invalid_fields = set()
        for cells in ws.iter_rows():
            for field in ("floor_height", "floors"):
                if any(cell.value == field for cell in cells):
                    ws.cell(cells[0].row, 3, "错误演示：不是数字")
                    invalid_fields.add(field)
        if len(invalid_fields) != 2:
            raise ValueError("无法定位 Excel 层高与层数字段")
        wb.save(invalid_workbook)
        # 固定压缩包和属性时间戳，重复导出的演示文件哈希保持稳定。
        for path in [workbook, invalid_workbook]:
            buffer = BytesIO()
            with (
                ZipFile(path) as source_zip,
                ZipFile(buffer, "w", ZIP_DEFLATED) as target_zip,
            ):
                for name in sorted(source_zip.namelist()):
                    content = source_zip.read(name)
                    if name == "docProps/core.xml":
                        content = re.sub(
                            rb"(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)",
                            rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                            content,
                        )
                    item = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
                    item.compress_type = ZIP_DEFLATED
                    target_zip.writestr(item, content)
            path.write_bytes(buffer.getvalue())
        imports = []
        for path in [
            workbook,
            invalid_workbook,
            ROOT / "data/schemes/方案录入模板.xlsx",
        ]:
            if not path.exists():
                continue
            content = path.read_bytes()
            parsed = parse_workbook(content)
            imports.append(
                {
                    "filename": path.name,
                    "sha256": hashlib.sha256(content).hexdigest(),
                    "schemes": parsed.schemes,
                    "project_params": parsed.project_params,
                    "errors": [
                        {
                            "scheme": e.scheme,
                            "field": e.field,
                            "label": e.label,
                            "message": e.message,
                        }
                        for e in parsed.errors
                    ],
                }
            )
        bundle = {
            "fixture_id": fixture_id,
            "disclaimer": "虚构数据，仅供开发调试，不得用于验证或答辩。评分与角色分析由真实引擎离线生成。",
            "source_hashes": hashes,
            "grid_step": args.grid_step,
            "projects": [project_read],
            "schemes": scheme_reads,
            "evaluation": evaluation,
            "roles": roles,
            "ranks": ranks,
            # 不带 session：提交进仓库的前端产物必须只含**内建基线**，
            # 使用者运行期加的自定义内容不能混进 git（见 app/api/schema.py 的说明）。
            "schema": build_all_schemas(),
            "rules": build_rules_response().model_dump(mode="json"),
            "imports": imports,
            "defaults": {
                "project": ProjectParams().model_dump(),
                "scheme": SchemeParams().model_dump(),
            },
            "validation": {
                "project": ProjectParams.model_json_schema(),
                "scheme": SchemeParams.model_json_schema(),
            },
        }
        (out / "data.json").write_text(
            json.dumps(bundle, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
            encoding="utf-8",
        )
        (out / "grid.json").write_text(
            json.dumps(grid, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
            + "\n",
            encoding="utf-8",
        )
    engine.dispose()
    print(
        f"已导出 {len(scheme_reads)} 个方案、{len(ranks)} 个角色、{len(grid)} 个精确网格快照：{out}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
