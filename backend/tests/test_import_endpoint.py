"""POST /api/projects/{id}/schemes/import 端点测试。

覆盖：
- 正常导入：方案全部成功，返回 200 + created 列表。
- 默认行为：xlsx 有错误，整个拒绝，返回 422 + errors。
- --allow-partial：跳过有错误的方案，返回 200 + errors。
- 空文件 / 文件类型错误。
- 项目不存在 → 404。
"""

from __future__ import annotations

import io

from app.excel import write_template
from fastapi.testclient import TestClient


def _create_project(client: TestClient, name: str = "导入测试项目") -> int:
    resp = client.post(
        "/api/projects",
        json={"name": name, "params": {"site_area": 18000}},
    )
    assert resp.status_code == 201
    return resp.json()["id"]


def _fill_template(client: TestClient, project_id: int, edits) -> bytes:
    """用 ``edits(workbook)`` 在内存中改模板，返回 xlsx 字节。"""
    import tempfile
    from pathlib import Path

    from openpyxl import load_workbook

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = Path(f.name)
    try:
        write_template(tmp_path)
        wb = load_workbook(tmp_path)
        edits(wb)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        tmp_path.unlink(missing_ok=True)


class TestImportEndpoint:
    def test_import_happy_path(self, client: TestClient) -> None:
        pid = _create_project(client)

        def edits(wb) -> None:
            # 项目信息
            name_col = next(
                c.column for c in wb["项目信息"][1] if c.value == "__name__"
            )
            for r in range(2, wb["项目信息"].max_row + 1):
                if wb["项目信息"].cell(row=r, column=name_col).value == "site_area":
                    wb["项目信息"].cell(row=r, column=2, value=20000)
            # 方案 A
            scheme_ws = wb["方案参数"]
            sname_col = next(
                c.column for c in scheme_ws[1] if c.value == "__name__"
            )
            for r in range(2, scheme_ws.max_row + 1):
                name = scheme_ws.cell(row=r, column=sname_col).value
                if name == "gfa":
                    scheme_ws.cell(row=r, column=3, value=15000)
                elif name == "structure_type":
                    scheme_ws.cell(row=r, column=3, value="剪力墙")

        data = _fill_template(client, pid, edits)
        resp = client.post(
            f"/api/projects/{pid}/schemes/import",
            files={"file": ("schemes.xlsx", data, "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert len(body["created"]) == 1
        assert body["errors"] == []
        # 实际写入数据库
        listing = client.get(f"/api/projects/{pid}/schemes").json()
        assert len(listing) == 1
        assert listing[0]["params"]["gfa"] == 15000
        assert listing[0]["params"]["structure_type"] == "剪力墙"

    def test_import_default_rejects_on_error(self, client: TestClient) -> None:
        """默认行为：xlsx 里有错误就拒绝整文件，不写数据库。"""
        pid = _create_project(client)

        def edits(wb) -> None:
            scheme_ws = wb["方案参数"]
            sname_col = next(
                c.column for c in scheme_ws[1] if c.value == "__name__"
            )
            for r in range(2, scheme_ws.max_row + 1):
                if scheme_ws.cell(row=r, column=sname_col).value == "floor_height":
                    scheme_ws.cell(row=r, column=3, value="三米")  # 数字格填文字

        data = _fill_template(client, pid, edits)
        resp = client.post(
            f"/api/projects/{pid}/schemes/import",
            files={"file": ("bad.xlsx", data, "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 422
        body = resp.json()
        assert "errors" in body["detail"]
        assert any("层高" in e["label"] for e in body["detail"]["errors"])

        # 数据库没有被写入
        listing = client.get(f"/api/projects/{pid}/schemes").json()
        assert listing == []

    def test_import_allow_partial(self, client: TestClient) -> None:
        pid = _create_project(client)

        def edits(wb) -> None:
            scheme_ws = wb["方案参数"]
            sname_col = next(
                c.column for c in scheme_ws[1] if c.value == "__name__"
            )
            # 方案A 正确，方案B 错误
            for r in range(2, scheme_ws.max_row + 1):
                name = scheme_ws.cell(row=r, column=sname_col).value
                if name == "gfa":
                    scheme_ws.cell(row=r, column=3, value=15000)  # A: OK
                    scheme_ws.cell(row=r, column=4, value="一万八")  # B: 错误
                elif name == "structure_type":
                    scheme_ws.cell(row=r, column=3, value="剪力墙")

        data = _fill_template(client, pid, edits)
        resp = client.post(
            f"/api/projects/{pid}/schemes/import?allow_partial=true",
            files={"file": ("partial.xlsx", data, "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert len(body["created"]) == 1
        assert body["created"][0]["name"] == "方案A"
        assert any("总建筑面积" in e["label"] for e in body["errors"])

        listing = client.get(f"/api/projects/{pid}/schemes").json()
        # 只有方案 A 被写入
        assert len(listing) == 1

    def test_import_empty_file_rejected(self, client: TestClient) -> None:
        pid = _create_project(client)
        resp = client.post(
            f"/api/projects/{pid}/schemes/import",
            files={"file": ("empty.xlsx", b"", "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 400

    def test_import_project_not_found(self, client: TestClient) -> None:
        resp = client.post(
            "/api/projects/9999/schemes/import",
            files={"file": ("x.xlsx", b"x", "application/octet-stream")},
        )
        assert resp.status_code == 404

    def test_import_duplicate_name_gets_suffix(self, client: TestClient) -> None:
        """项目内已有同名方案时，新方案自动重命名。"""
        pid = _create_project(client)
        # 先建一个"方案A"
        client.post(
            f"/api/projects/{pid}/schemes",
            json={"name": "方案A", "params": {"gfa": 10000}},
        )

        def edits(wb) -> None:
            ws = wb["方案参数"]
            sname_col = next(c.column for c in ws[1] if c.value == "__name__")
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=sname_col).value == "gfa":
                    ws.cell(row=r, column=3, value=15000)

        data = _fill_template(client, pid, edits)
        resp = client.post(
            f"/api/projects/{pid}/schemes/import",
            files={"file": ("dup.xlsx", data, "application/vnd.openxmlformats")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["created"][0]["name"] == "方案A (2)"

        listing = client.get(f"/api/projects/{pid}/schemes").json()
        names = sorted(s["name"] for s in listing)
        assert names == ["方案A", "方案A (2)"]