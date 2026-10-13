"""动态模板、导出及严格导入必须采用同一运行期字段注册表。"""

from io import BytesIO

from app.db import init_db, override_engine_for_testing
from app.excel import parse_workbook, write_export
from app.models import EvalResultRow, Project, Scheme
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import StaticPool, inspect
from sqlmodel import Session, SQLModel, create_engine


def _field(client: TestClient) -> None:
    response = client.post("/api/custom/fields", json={
        "key": "demo_ratio", "label": "演示比例", "kind": "ratio", "group": "基本信息",
    })
    assert response.status_code == 201, response.text


def test_dynamic_template_and_export_roundtrip(client, engine):
    _field(client)
    template = client.get("/api/excel/template")
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content))
    assert any(cell.value == "demo_ratio" for row in workbook["方案参数"] for cell in row)
    assert not any(cell.value == "custom" for row in workbook["方案参数"] for cell in row)
    project = client.post("/api/projects", json={"name": "Excel测试"}).json()
    created = client.post(f"/api/projects/{project['id']}/schemes", json={
        "name": "零值", "params": {"custom": {"demo_ratio": 0}},
    })
    assert created.status_code == 201, created.text
    exported = client.get(f"/api/excel/projects/{project['id']}/export")
    assert exported.status_code == 200
    with Session(engine) as session:
        parsed = parse_workbook(exported.content, session=session)
    assert not parsed.errors
    assert parsed.schemes[0]["params"]["custom"]["demo_ratio"] == 0
    assert client.get("/api/excel/projects/999999/export").status_code == 404


def test_invalid_custom_excel_rejects_entire_import(client, engine, tmp_path):
    _field(client)
    with Session(engine) as session:
        path = write_export(tmp_path / "invalid.xlsx", {}, [
            {"name": "合法", "params": {"custom": {"demo_ratio": 0.5}}},
            {"name": "不合法", "params": {"custom": {"demo_ratio": 2}}},
        ], session=session)
    project = client.post("/api/projects", json={"name": "原子导入"}).json()
    response = client.post(f"/api/projects/{project['id']}/schemes/import", files={
        "file": ("invalid.xlsx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    })
    assert response.status_code == 422, response.text
    errors = response.json()["detail"]["errors"]
    assert any(error["scheme"] == "不合法" and error["field"] == "demo_ratio" and error["label"] == "演示比例" for error in errors)
    assert client.get(f"/api/projects/{project['id']}/schemes").json() == []


def test_legacy_database_adds_custom_tables_without_rewriting_inputs(engine):
    legacy = create_engine("sqlite://", poolclass=StaticPool)
    SQLModel.metadata.create_all(legacy, tables=[
        Project.__table__, Scheme.__table__, EvalResultRow.__table__,
    ])
    original = {"floors": 17, "floor_height": 3.0}
    with Session(legacy) as session:
        project = Project(name="旧项目", params={"site_area": 1000})
        session.add(project)
        session.flush()
        scheme = Scheme(project_id=project.id, name="旧方案", params=original)
        session.add(scheme)
        session.commit()
        scheme_id = scheme.id
    try:
        override_engine_for_testing(legacy)
        init_db()
        assert "custom_fields" in inspect(legacy).get_table_names()
        assert "custom_rules" in inspect(legacy).get_table_names()
        with Session(legacy) as session:
            assert session.get(Scheme, scheme_id).params == original
    finally:
        override_engine_for_testing(engine)
        legacy.dispose()


def test_export_formula_looking_user_text_stays_literal(client, engine):
    registered = client.post("/api/custom/fields", json={
        "key": "formula_text", "label": "=1+1", "kind": "text", "group": "基本信息",
    })
    assert registered.status_code == 201
    project = client.post("/api/projects", json={"name": "原文往返"}).json()
    value = '=HYPERLINK("https://example.invalid","示例")'
    created = client.post(f"/api/projects/{project['id']}/schemes", json={
        "name": "=1+1", "params": {"custom": {"formula_text": value}},
    })
    assert created.status_code == 201
    exported = client.get(f"/api/excel/projects/{project['id']}/export")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    with Session(engine) as session:
        parsed = parse_workbook(exported.content, session=session)
    assert not parsed.errors
    assert parsed.schemes[0]["name"] == "=1+1"
    assert parsed.schemes[0]["params"]["custom"]["formula_text"] == value
