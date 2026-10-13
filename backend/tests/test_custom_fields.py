"""自定义字段（覆盖层第二片）：注册表校验、并入 schema、摊平进求值命名空间。

守住五条：
1. 只能**追加**：字段名不得与内建字段、派生量或保留字段重名；
2. 只有方案层支持自定义字段（值存在 ``SchemeParams.custom`` 里）；
3. ``GET /api/schema`` 会把它们追加成独立条目（``"custom": true``），
   于是 schema 驱动的表单与 Excel 列自动带上；
4. 引擎构造命名空间时把 ``custom`` 的键**摊平到顶层**，自定义规则可以直接引用；
5. **内建规则不能引用自定义字段**——``check_rules`` 是纯文件校验、看不到覆盖层，
   所以校验内建规则时"已知字段"里不该有它们（这里用一条测试把这条界线钉住）。
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from app.config import get_settings
from app.custom import fields as fields_service
from app.engine.rules import (
    build_namespace,
    known_field_names,
    load_rules_from_dir,
    set_ruleset_for_testing,
)
from app.models import ProjectParams, SchemeParams
from fastapi.testclient import TestClient

GOOD_RULES = Path(__file__).resolve().parent / "fixtures" / "rules_good"


@pytest.fixture(autouse=True)
def fixed_ruleset() -> Iterator[None]:
    """整个模块用固定的测试规则集，不受 ``rules/`` 里种子规则变化影响。"""
    set_ruleset_for_testing(load_rules_from_dir(GOOD_RULES))
    yield
    set_ruleset_for_testing(None)


FIELD = {
    "key": "bearing_wall_ratio",
    "label": "承重墙占比",
    "kind": "ratio",
    "unit": "—",
    "group": "结构与造价",
    "description": "本次比选自定指标：承重墙占外墙比例",
}


def _create(client: TestClient, **overrides: Any) -> dict[str, Any]:
    response = client.post("/api/custom/fields", json={**FIELD, **overrides})
    assert response.status_code == 201, response.text
    return response.json()


def _scheme_schema(client: TestClient) -> dict[str, Any]:
    return client.get("/api/schema/scheme").json()


def test_added_field_shows_up_in_schema_and_can_be_removed(client: TestClient) -> None:
    before = _scheme_schema(client)
    assert "bearing_wall_ratio" not in {f["name"] for f in before["fields"]}

    created = _create(client)
    assert created["key"] == "bearing_wall_ratio"
    assert created["kind"] == "ratio"

    after = _scheme_schema(client)
    entry = next(f for f in after["fields"] if f["name"] == "bearing_wall_ratio")
    assert entry["label"] == "承重墙占比"
    assert entry["group"] == "结构与造价"
    assert entry["custom"] is True
    assert entry["constraints"] == {"minimum": 0, "maximum": 1}
    # 内建字段一条不少，且仍排在前面
    assert [f["name"] for f in before["fields"]] == [
        f["name"] for f in after["fields"][: len(before["fields"])]
    ]

    assert client.delete(f"/api/custom/fields/{created['id']}").status_code == 204
    assert "bearing_wall_ratio" not in {f["name"] for f in _scheme_schema(client)["fields"]}


def test_key_must_not_shadow_builtin_fields(client: TestClient) -> None:
    reserved = client.post("/api/custom/fields", json={**FIELD, "key": "floor_height"})
    assert reserved.status_code == 422
    problem = reserved.json()["detail"]["problems"][0]
    assert problem["code"] == "FIELD_KEY_RESERVED"
    assert "只能**追加**" in problem["fix"] or "追加" in problem["fix"]

    derived = client.post("/api/custom/fields", json={**FIELD, "key": "shape_coef"})
    assert derived.status_code == 422
    assert derived.json()["detail"]["problems"][0]["code"] == "FIELD_KEY_RESERVED"

    container = client.post("/api/custom/fields", json={**FIELD, "key": "custom"})
    assert container.status_code == 422


def test_key_label_kind_and_group_are_validated(client: TestClient) -> None:
    bad_key = client.post("/api/custom/fields", json={**FIELD, "key": "承重墙"})
    assert bad_key.status_code == 422
    assert bad_key.json()["detail"]["problems"][0]["code"] == "FIELD_KEY_SHAPE"

    blank_label = client.post("/api/custom/fields", json={**FIELD, "label": "   "})
    assert blank_label.status_code == 422
    assert blank_label.json()["detail"]["problems"][0]["code"] == "FIELD_LABEL_EMPTY"

    bad_kind = client.post("/api/custom/fields", json={**FIELD, "kind": "颜色"})
    assert bad_kind.status_code == 422
    assert bad_kind.json()["detail"]["problems"][0]["code"] == "FIELD_KIND_UNKNOWN"

    bad_group = client.post("/api/custom/fields", json={**FIELD, "group": "随便编的"})
    assert bad_group.status_code == 422
    assert bad_group.json()["detail"]["problems"][0]["code"] == "FIELD_GROUP_UNKNOWN"

    options_on_number = client.post(
        "/api/custom/fields", json={**FIELD, "options": ["甲", "乙"]}
    )
    assert options_on_number.status_code == 422
    assert (
        options_on_number.json()["detail"]["problems"][0]["code"] == "FIELD_OPTIONS_UNSUPPORTED"
    )

    _create(client)
    duplicate = client.post("/api/custom/fields", json=FIELD)
    assert duplicate.status_code == 422
    assert duplicate.json()["detail"]["problems"][0]["code"] == "FIELD_KEY_DUPLICATE"

    missing = client.delete("/api/custom/fields/9999")
    assert missing.status_code == 422
    assert missing.json()["detail"]["problems"][0]["code"] == "FIELD_NOT_FOUND"


def test_custom_values_are_flattened_into_the_namespace() -> None:
    """引擎把 custom 的键摊平到顶层；内建字段优先，不会被顶掉。"""
    project = ProjectParams(site_area=18000, far_max=2.5, density_max=0.22, height_max=54)
    scheme = SchemeParams(
        structure_type="剪力墙",
        floors=18,
        floor_height=3.0,
        custom={"bearing_wall_ratio": 0.62, "floor_height": 999.0},
    )
    namespace = build_namespace(project, scheme)
    assert namespace.values["bearing_wall_ratio"] == 0.62
    # 内建字段永远优先：custom 里同名（不该发生，注册表已拦）也不能顶掉
    assert namespace.values["floor_height"] == 3.0

    # 没填的自定义字段按"缺数据"处理，而不是判成不通过
    empty = build_namespace(project, SchemeParams(structure_type="剪力墙", custom={}))
    assert "bearing_wall_ratio" in empty.missing_for(["bearing_wall_ratio"])


def test_custom_fields_are_known_to_custom_rules_only(client: TestClient) -> None:
    """自定义字段进自定义规则的"已知字段"，但不进内建规则的。"""
    assert "bearing_wall_ratio" not in known_field_names()

    _create(client)
    from app.custom.rules import add_rule, ruleset_with_custom
    from app.db import get_engine
    from sqlmodel import Session

    with Session(get_engine()) as session:
        # 校验器认识它了（自定义规则的表达式可以引用）
        assert "bearing_wall_ratio" in fields_service.known_names(session)

        rule = {
            "id": "X-07",
            "name": "自定义：承重墙占比不宜过低",
            "type": "metric",
            "dimension": "structure",
            "source": "经验规则（结构组）：本期自定，承重墙占比低则抗侧刚度不足",
            "value": "bearing_wall_ratio",
            "scoring": {"kind": "piecewise", "points": [[0.3, 0], [0.6, 100]]},
        }
        row = add_rule(session, payload=rule)
        assert row.rule_id == "X-07"
        merged = ruleset_with_custom(session)
        assert merged.by_id("X-07") is not None

        # 反过来：内建规则仍不认识它（check_rules 是文件校验，看不到覆盖层）
        from app.engine.rule_check import check_rule_problems

        problems = check_rule_problems(
            {
                "id": "S-99",
                "name": "内建规则引用自定义字段",
                "type": "hard",
                "source": "测试",
                "check": "bearing_wall_ratio >= 0.5",
            }
        )
        assert [p.code for p in problems] == ["BAD_EXPRESSION"]
        # 但把覆盖层的字段名传进去就通过了 —— 差别只在这份"已知字段"清单
        assert (
            check_rule_problems(
                {
                    "id": "S-99",
                    "name": "自定义规则引用自定义字段",
                    "type": "hard",
                    "source": "测试",
                    "check": "bearing_wall_ratio >= 0.5",
                },
                known=fields_service.known_names(session),
            )
            == []
        )


def test_writes_are_closed_outside_dev(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("APP_ENV", "prod")
    get_settings.cache_clear()
    try:
        blocked = client.post("/api/custom/fields", json=FIELD)
        assert blocked.status_code == 403, blocked.text
        assert client.get("/api/custom/fields").status_code == 200
        assert client.get("/api/schema/scheme").status_code == 200
    finally:
        get_settings.cache_clear()


def test_custom_field_round_trips_through_excel(
    client: TestClient, tmp_path: Path
) -> None:
    """自定义字段能进 Excel 列、能读回来、并且落在 ``params["custom"]`` 里。

    这是"界面上能填的字段 == Excel 里能填的列"那条一致性的实测：
    模板生成、导出、解析三条路径都走 ``app/schemas/merged.py`` 的同一份字段表。
    """
    _create(client)

    from app.db import get_engine
    from app.excel import SHEET_SCHEMES, parse_workbook, write_export, write_template
    from openpyxl import load_workbook
    from sqlmodel import Session

    with Session(get_engine()) as session:
        # 1) 模板里出现了这个字段的中文名（列被加上了）
        template = write_template(tmp_path / "tpl.xlsx", session=session)
        sheet = load_workbook(template)[SHEET_SCHEMES]
        labels = [sheet.cell(row=r, column=1).value for r in range(1, sheet.max_row + 1)]
        assert FIELD["label"] in labels

        # 2) 导出 → 解析：值从 custom 字典来，又回到 custom 字典去
        exported = write_export(
            tmp_path / "out.xlsx",
            {
                "site_area": 20000,
                "far_max": 2.5,
                "density_max": 0.3,
                "height_max": 80,
                "climate_zone": "夏热冬冷",
            },
            [
                {
                    "name": "A 方案",
                    "params": {
                        "floor_height": 3.1,
                        "custom": {"bearing_wall_ratio": 0.62},
                    },
                }
            ],
            session=session,
        )
        result = parse_workbook(exported, session=session)

    assert result.errors == []
    params = result.schemes[0]["params"]
    assert params["floor_height"] == pytest.approx(3.1)
    # 落在 custom 里，而不是顶层——顶层会被 SchemeParams 丢掉
    assert params["custom"] == {"bearing_wall_ratio": pytest.approx(0.62)}

    # 3) 不给 session 时是纯内建基线：列不出现（构建产物只含内建字段）
    plain = write_template(tmp_path / "builtin.xlsx")
    plain_sheet = load_workbook(plain)[SHEET_SCHEMES]
    plain_labels = [
        plain_sheet.cell(row=r, column=1).value for r in range(1, plain_sheet.max_row + 1)
    ]
    assert FIELD["label"] not in plain_labels
